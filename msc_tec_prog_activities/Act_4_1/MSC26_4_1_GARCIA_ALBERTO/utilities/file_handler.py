import csv

from openpyxl import Workbook, load_workbook

#Lectura y escritura de archivos CSV / XLSX.
#Solo trabaja con listas de filas (listas de valores), no sabe nada de cuentas ni clientes;
#la conversión entre objetos y filas la hace la ventana principal.


#--------------------------------------------------------------------------
# CSV
#--------------------------------------------------------------------------
def write_csv(path, headers, rows):
    #newline="" evita que en Windows se agregue una línea vacía entre filas
    #utf-8-sig agrega la marca BOM para que Excel muestre bien los acentos
    with open(path, "w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file)
        writer.writerow(headers)
        writer.writerows(rows)


def read_csv(path):
    #utf-8-sig también lee archivos sin BOM, así que sirve para ambos casos
    with open(path, "r", newline="", encoding="utf-8-sig") as file:
        reader = csv.reader(file)
        next(reader, None) #Se salta el encabezado
        return list(reader)


#--------------------------------------------------------------------------
# XLSX
#--------------------------------------------------------------------------
def write_xlsx(path, headers, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cuentas"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def read_xlsx(path):
    #data_only=True regresa el valor calculado de las fórmulas en lugar de la fórmula
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    #min_row=2 se salta el encabezado; values_only regresa los valores y no las celdas
    #Las celdas vacías llegan como None
    return [list(row) for row in sheet.iter_rows(min_row=2, values_only=True)]


#--------------------------------------------------------------------------
# SEGÚN LA EXTENSIÓN
#--------------------------------------------------------------------------
def write_file(path, headers, rows):
    if path.lower().endswith(".xlsx"):
        write_xlsx(path, headers, rows)
    elif path.lower().endswith(".csv"):
        write_csv(path, headers, rows)
    else:
        raise ValueError("Formato no soportado, usa .csv o .xlsx")


def read_file(path):
    if path.lower().endswith(".xlsx"):
        return read_xlsx(path)
    if path.lower().endswith(".csv"):
        return read_csv(path)
    raise ValueError("Formato no soportado, usa .csv o .xlsx")
